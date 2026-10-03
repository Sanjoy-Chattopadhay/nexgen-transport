"""Trips: every geofence a trip touched, and how far to trust each crossing.

The trip page is the geofence half of a trip's story. Commercial detail
(invoice, ETA, consignee) belongs to the fleet system; what this adds is the
sequence of places the truck was actually at, how long at each, how it moved
between them, where the GPS could not see it, and what that means for the
fleet system's own timestamps.
"""

from __future__ import annotations

from datetime import timedelta

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response

from nexgen.shared.geoengine.api.v1.common import (FACILITY, clean, day_bounds, order_clause, paged, resolve_run,
                                      rows, scale_sql)
from nexgen.shared.geoengine.db import get_geo_db
from nexgen.shared.geoengine.osrm import polyline
from nexgen.shared.geoengine.pipeline import phases
from nexgen.shared.geoengine.prep import codec

router = APIRouter(tags=["trips"])

TRIP_SORT = {
    "start": "s.dt_first_ping", "end": "s.dt_last_ping", "trip": "s.i_trip_no",
    "visits": "s.i_facility_visits", "places": "s.i_places", "dwell": "s.i_facility_dwell_s",
    "violations": "s.i_violations", "quality": "s.s_quality", "distance": "s.d_distance_km",
    "transit": "s.i_transit_s", "vehicle": "s_asset_id", "transporter": "m.s_trans_name",
    "pings": "s.i_pings_read", "loading": "p.i_loading_s", "unloading": "p.i_unloading_s",
}

TRIP_COLS = """s.i_trip_no, COALESCE(m.s_asset_id, s.s_asset_id) AS s_asset_id, m.s_asset_type,
    m.s_driver_name, m.s_driver_mobile, m.s_trans_name, m.s_cnr_name, m.s_cne_name,
    m.s_origin, m.s_destination, m.s_status, m.s_trip_class, m.dt_booking, m.dt_trip_start,
    m.dt_trip_ata, m.dt_trip_end, s.dt_first_ping, s.dt_last_ping, s.i_pings_read, s.i_pings_used,
    s.i_pings_dropped, s.s_quality, s.s_quality_reason, s.i_visits, s.i_facility_visits,
    s.i_places, s.i_facility_dwell_s, s.i_violations, s.i_inferred, s.i_spikes, s.i_medians,
    s.i_snapped, s.i_stops, s.i_moving_gaps, s.i_moving_gap_s, s.d_distance_km,
    s.d_gap_distance_km, s.s_osrm, s.i_first_site_id, s.s_first_site, s.dt_first_enter,
    s.dt_first_exit, s.i_last_site_id, s.s_last_site, s.dt_last_enter, s.dt_last_exit,
    s.i_transit_s, s.i_max_gap_seconds, sh.i_siblings, sh.s_sibling_trips"""

# Every query selecting TRIP_COLS joins these.
TRIP_FROM = """FROM geo_trip_summary s
    LEFT JOIN geo_trip_meta m ON m.i_trip_no = s.i_trip_no
    LEFT JOIN geo_trip_share sh ON sh.i_run_id = s.i_run_id AND sh.i_trip_no = s.i_trip_no"""

# The trip as a number line (pipeline/phases.py), for the list's small bar.
PHASE_COLS = """p.s_shape, p.i_loading_s, p.i_unloading_s, p.i_transit_moving_s, p.i_transit_stop_s,
    p.i_transit_halt_s, p.i_transit_silent_s, p.i_span_s, p.j_bar"""
PHASE_JOIN = "LEFT JOIN geo_trip_phase p ON p.i_run_id = s.i_run_id AND p.i_trip_no = s.i_trip_no"


def _trip_filters(rid, q, transporter, vehicle, driver, status, quality, date_from, date_to,
                  has_alerts, origin, destination, site_id):
    where, params = ["s.i_run_id = %s"], [rid]
    if q:
        if q.isdigit():
            where.append("s.i_trip_no = %s")
            params.append(int(q))
        else:
            where.append("(COALESCE(m.s_asset_id, s.s_asset_id) LIKE %s OR m.s_driver_name LIKE %s "
                         "OR m.s_trans_name LIKE %s OR m.s_origin LIKE %s OR m.s_destination LIKE %s)")
            params += [f"%{q}%"] * 5
    for col, val in (("m.s_trans_name", transporter), ("m.s_driver_name", driver),
                     ("m.s_status", status), ("s.s_quality", quality),
                     ("m.s_origin", origin), ("m.s_destination", destination)):
        if val:
            where.append(f"{col} = %s")
            params.append(val)
    if vehicle:
        where.append("COALESCE(m.s_asset_id, s.s_asset_id) = %s")
        params.append(vehicle)
    a, b = day_bounds(date_from, date_to)
    if a:
        where.append("s.dt_last_ping >= %s")
        params.append(a)
    if b:
        where.append("s.dt_first_ping < %s")
        params.append(b)
    if has_alerts is True:
        where.append("s.i_violations > 0")
    if site_id:
        where.append("EXISTS (SELECT 1 FROM geo_visit v WHERE v.i_run_id = s.i_run_id "
                     "AND v.i_trip_no = s.i_trip_no AND v.i_site_id = %s)")
        params.append(site_id)
    return where, params


@router.get("/trips")
def list_trips(
    q: str | None = None, transporter: str | None = None, vehicle: str | None = None,
    driver: str | None = None, status: str | None = None, quality: str | None = None,
    date_from: str | None = Query(None, alias="from"), date_to: str | None = Query(None, alias="to"),
    has_alerts: bool | None = None, origin: str | None = None, destination: str | None = None,
    site_id: int | None = None,
    sort: str = "start", order: str = "desc",
    page: int = Query(1, ge=1), page_size: int = Query(25, ge=1, le=500),
    run: int | None = None, conn=Depends(get_geo_db),
):
    with conn.cursor() as cur:
        rid = resolve_run(cur, run)
        where, params = _trip_filters(rid, q, transporter, vehicle, driver, status, quality,
                                      date_from, date_to, has_alerts, origin, destination, site_id)
        base = f"""{TRIP_FROM}
                   WHERE {' AND '.join(where)}"""
        # Totals sum each trip's share of the physical ledger, so a stay seen by
        # three consignments on one truck counts once (pipeline/physical.py);
        # a trip without a share yet counts its own figure.
        cur.execute(f"""SELECT COUNT(*) trips,
                               SUM(COALESCE(sh.i_facility_visits, s.i_facility_visits)) facility_visits,
                               SUM(COALESCE(sh.i_facility_dwell_s, s.i_facility_dwell_s)) facility_dwell_s,
                               SUM(COALESCE(sh.i_alerts, s.i_violations)) violations,
                               SUM(COALESCE(sh.i_inferred, s.i_inferred)) inferred,
                               ROUND(SUM(COALESCE(sh.d_distance_km, s.d_distance_km)), 1) distance_km,
                               SUM(sh.i_siblings > 0) sharing_gps,
                               SUM(s.s_quality = 'good') good, SUM(s.s_quality = 'broken') broken,
                               SUM(s.i_places >= 2) with_lane,
                               COUNT(DISTINCT COALESCE(m.s_asset_id, s.s_asset_id)) vehicles
                          {base}""", params)
        kpis = clean(cur.fetchone())
        cur.execute(f"""SELECT {TRIP_COLS}, {PHASE_COLS} {TRIP_FROM} {PHASE_JOIN}
                         WHERE {' AND '.join(where)}
                        {order_clause(sort, order, TRIP_SORT, 'start')}, s.i_trip_no DESC
                        LIMIT %s OFFSET %s""", (*params, page_size, (page - 1) * page_size))
        items = rows(cur)
        cur.execute("SELECT DISTINCT s_status FROM geo_trip_meta WHERE s_status IS NOT NULL ORDER BY 1")
        statuses = [r["s_status"] for r in cur.fetchall()]
    return {"run_id": rid, "kpis": kpis, "statuses": statuses,
            "qualities": ["good", "sparse", "noisy", "broken", "no_gps"],
            **paged(kpis["trips"], page, page_size, items)}


def _fit_mix(records: list[dict]) -> list[dict]:
    mix: dict[tuple, dict] = {}
    for r in records:
        key = (r["fit"], r["role"])
        m = mix.setdefault(key, {"s_fit": r["fit"], "s_role": r["role"], "n": 0, "max_shift_m": 0.0})
        m["n"] += 1
        if r["shift_m"] is not None:
            m["max_shift_m"] = max(m["max_shift_m"], round(r["shift_m"]))
    return sorted(mix.values(), key=lambda m: -m["n"])


def _places(visits: list[dict]) -> list[dict]:
    """Nested facility visits collapsed into places; see pipeline/runner.places."""
    fac = sorted((v for v in visits if v["s_scale"] in FACILITY), key=lambda v: v["dt_enter"])
    clusters: list[dict] = []
    for v in fac:
        end = v["dt_exit"] or v["dt_enter"] + timedelta(seconds=v["i_dwell_seconds"] or 0)
        if clusters and v["dt_enter"] <= clusters[-1]["end"]:
            c = clusters[-1]
            c["members"].append(v)
            c["end"] = max(c["end"], end)
            c["open"] = c["open"] or bool(v["b_open"])
        else:
            clusters.append({"start": v["dt_enter"], "end": end, "open": bool(v["b_open"]),
                             "members": [v]})
    out = []
    for c in clusters:
        outer = max(c["members"], key=lambda v: (v["d_area_sqm"] or 0, -v["i_site_id"]))
        inner = min(c["members"], key=lambda v: (v["d_area_sqm"] or 0, v["i_site_id"]))
        out.append({
            "site_id": outer["i_site_id"], "name": outer["s_site_name"], "scale": outer["s_scale"],
            "category": outer["s_category"],
            "innermost_site_id": inner["i_site_id"], "innermost": inner["s_site_name"],
            "enter": c["start"], "exit": None if c["open"] else c["end"],
            "last_seen": c["end"], "open": c["open"],
            "dwell_s": int((c["end"] - c["start"]).total_seconds()),
            "entry_observed": all(bool(m["b_entry_observed"]) for m in c["members"]
                                  if m["dt_enter"] == c["start"]),
            "zones": sorted({m["s_site_name"] for m in c["members"]}),
            "enter_gap_s": min((m["i_enter_gap_seconds"] for m in c["members"]
                                if m["dt_enter"] == c["start"] and m["i_enter_gap_seconds"] is not None),
                               default=None),
        })
    return out


@router.get("/trips/{trip_no}")
def trip_detail(trip_no: int, run: int | None = None, conn=Depends(get_geo_db)):
    with conn.cursor() as cur:
        rid = resolve_run(cur, run)
        cur.execute(f"""SELECT {TRIP_COLS}, s.i_inside_seconds, s.d_coverage_pct, s.i_distinct_sites
                          {TRIP_FROM}
                         WHERE s.i_run_id=%s AND s.i_trip_no=%s""", (rid, trip_no))
        trip = cur.fetchone()
        if not trip:
            raise HTTPException(404, f"trip {trip_no} is not in run {rid}")
        trip = clean(trip)

        cur.execute(f"""
            SELECT v.id, v.i_site_id, v.s_site_name, v.s_type, v.s_category, v.s_scale, v.dt_enter,
                   v.dt_exit, v.b_open, v.b_primary, v.b_entry_observed, v.i_dwell_seconds, v.i_pings,
                   v.i_max_speed, ROUND(v.d_distance_m) d_distance_m, v.i_enter_gap_seconds,
                   v.i_exit_gap_seconds, v.s_confirmed_by, f.d_area_sqm, f.d_inradius_m
              FROM geo_visit v LEFT JOIN geo_fence f ON f.i_fence_id = v.i_fence_id
             WHERE v.i_run_id=%s AND v.i_trip_no=%s
             ORDER BY v.dt_enter, f.d_area_sqm DESC""", (rid, trip_no))
        visits = rows(cur)

        cur.execute("""SELECT s_event, i_site_id, s_site_name, dt_event, i_gap_seconds, d_lat, d_long,
                              i_speed, s_confirmed_by
                         FROM geo_event WHERE i_run_id=%s AND i_trip_no=%s ORDER BY dt_event""",
                    (rid, trip_no))
        events = rows(cur)
        cur.execute("""SELECT s_kind, i_site_id, s_site_name, dt_event, i_observed, i_limit, s_detail,
                              d_lat, d_long
                         FROM geo_violation WHERE i_run_id=%s AND i_trip_no=%s ORDER BY dt_event""",
                    (rid, trip_no))
        violations = rows(cur)
        cur.execute("""SELECT i_seq, dt_start, dt_end, i_duration_s, i_pings, i_spikes, i_max_gap_s,
                              d_lat, d_long, d_p90_spread_m, i_site_id, s_site_name, s_scale
                         FROM geo_stop WHERE i_run_id=%s AND i_trip_no=%s ORDER BY dt_start""",
                    (rid, trip_no))
        stops = rows(cur)
        cur.execute("""SELECT id, dt_from, dt_to, i_gap_s, d_from_lat, d_from_long, d_to_lat, d_to_long,
                              d_straight_m, s_kind, s_route, d_route_m, d_route_s, i_unexplained_s
                         FROM geo_gap WHERE i_run_id=%s AND i_trip_no=%s ORDER BY dt_from""",
                    (rid, trip_no))
        gaps = rows(cur)
        cur.execute("""SELECT i_site_id, s_site_name, s_scale, s_category, s_kind, dt_gap_from, dt_gap_to,
                              dt_est_enter, dt_est_exit, d_inside_m, d_min_dist_m, s_confidence
                         FROM geo_inferred_visit WHERE i_run_id=%s AND i_trip_no=%s ORDER BY dt_gap_from""",
                    (rid, trip_no))
        inferred = rows(cur)
        cur.execute("""SELECT s_reason, i_count FROM geo_ping_reject WHERE i_run_id=%s AND i_trip_no=%s""",
                    (rid, trip_no))
        rejects = rows(cur)
        cur.execute("SELECT m_data FROM geo_fit_trail WHERE i_run_id=%s AND i_trip_no=%s",
                    (rid, trip_no))
        blob = cur.fetchone()
        fit_mix = _fit_mix(codec.records(blob["m_data"])) if blob else []
        cur.execute("SELECT s_variant, s_osrm, j_params FROM geo_run WHERE i_run_id=%s", (rid,))
        run_row = cur.fetchone()

    places = _places(visits)
    timeline = []
    for i, p in enumerate(places):
        timeline.append({"kind": "place", **p})
        if i + 1 < len(places):
            nxt = places[i + 1]
            leave = p["exit"] or p["last_seen"]
            legs_stops = [s for s in stops if s["i_site_id"] is None
                          and leave <= s["dt_start"] <= nxt["enter"]]
            legs_gaps = [g for g in gaps if g["s_kind"] == "moving"
                         and leave <= g["dt_from"] <= nxt["enter"]]
            timeline.append({
                "kind": "travel", "from": p["name"], "to": nxt["name"],
                "depart": p["exit"], "arrive": nxt["enter"],
                "travel_s": int((nxt["enter"] - leave).total_seconds()) if p["exit"] else None,
                "stops_outside": len(legs_stops),
                "stopped_s": sum(s["i_duration_s"] for s in legs_stops),
                "moving_gaps": len(legs_gaps),
                "unobserved_s": sum(g["i_gap_s"] for g in legs_gaps),
            })

    # The fleet system's own stamps against what the geofences saw.
    tms = None
    if trip.get("dt_trip_start") or trip.get("dt_booking"):
        first = places[0] if places else None
        last = places[-1] if len(places) > 1 else None

        def delta(a, b):
            return int((a - b).total_seconds()) if a and b else None
        tms = {
            "booking": trip.get("dt_booking"), "gate_out": trip.get("dt_trip_start"),
            "ata": trip.get("dt_trip_ata"),
            "first_place": first["name"] if first else None,
            "first_enter": first["enter"] if first else None,
            "first_exit": first["exit"] if first else None,
            # Positive: the truck physically left after the gate-out stamp.
            "exit_after_gate_out_s": delta(first["exit"] if first else None, trip.get("dt_trip_start")),
            "entry_after_booking_s": delta(first["enter"] if first else None, trip.get("dt_booking")),
            "last_place": last["name"] if last else None,
            "last_enter": last["enter"] if last else None,
            "arrival_vs_ata_s": delta(last["enter"] if last else None, trip.get("dt_trip_ata")),
        }

    return {"run_id": rid, "variant": run_row["s_variant"], "osrm": run_row["s_osrm"],
            "trip": trip, "places": places, "timeline": timeline, "visits": visits,
            "events": events, "violations": violations, "stops": stops, "gaps": gaps,
            "inferred": inferred, "rejects": rejects, "fit_mix": fit_mix, "tms": tms}


@router.get("/trips/{trip_no}/track")
def trip_track(trip_no: int, max_points: int = Query(6000, ge=100, le=50000),
               run: int | None = None, conn=Depends(get_geo_db)):
    """Raw and fitted positions side by side, plus gap routes and fence rings.

    Downsampled for drawing, but never by dropping the fixes that explain the
    result: spikes, refused fixes and stop boundaries are always kept.
    """
    with conn.cursor() as cur:
        rid = resolve_run(cur, run)
        cur.execute("""SELECT id, dt_message, d_lat, d_long, i_speed FROM geo_gps_ping
                        WHERE i_trip_no=%s ORDER BY dt_message, id""", (trip_no,))
        raw = cur.fetchall()
        if not raw:
            raise HTTPException(404, f"no GPS for trip {trip_no}")
        cur.execute("SELECT m_data FROM geo_fit_trail WHERE i_run_id=%s AND i_trip_no=%s",
                    (rid, trip_no))
        blob = cur.fetchone()
        fit = {r["id"]: r for r in codec.records(blob["m_data"])} if blob else {}
        cur.execute("""SELECT dt_from, dt_to, i_gap_s, s_kind, s_route, s_geometry, d_from_lat, d_from_long,
                              d_to_lat, d_to_long, i_unexplained_s
                         FROM geo_gap WHERE i_run_id=%s AND i_trip_no=%s AND s_kind='moving'""",
                    (rid, trip_no))
        gaps = cur.fetchall()
        cur.execute("""SELECT DISTINCT i_site_id FROM geo_visit WHERE i_run_id=%s AND i_trip_no=%s
                       UNION SELECT DISTINCT i_site_id FROM geo_inferred_visit WHERE i_run_id=%s AND i_trip_no=%s
                       UNION SELECT DISTINCT i_site_id FROM geo_stop
                             WHERE i_run_id=%s AND i_trip_no=%s AND i_site_id IS NOT NULL""",
                    (rid, trip_no, rid, trip_no, rid, trip_no))
        site_ids = [r["i_site_id"] for r in cur.fetchall()]
        fences = []
        if site_ids:
            marks = ",".join(["%s"] * len(site_ids))
            cur.execute(f"""SELECT i_site_id, s_site_name, s_category, {scale_sql('d_area_sqm')} s_scale
                              FROM geo_fence WHERE i_site_id IN ({marks})
                             ORDER BY d_area_sqm DESC LIMIT 150""", site_ids)
            meta = {r["i_site_id"]: r for r in cur.fetchall()}
            if meta:
                marks = ",".join(["%s"] * len(meta))
                cur.execute(f"""SELECT i_site_id, d_lat, d_long FROM geo_site_vertex
                                 WHERE i_site_id IN ({marks}) ORDER BY i_site_id, i_seq""", list(meta))
                rings: dict[int, list] = {}
                for r in cur.fetchall():
                    rings.setdefault(r["i_site_id"], []).append([float(r["d_lat"]), float(r["d_long"])])
                fences = [{"site_id": sid, "name": m["s_site_name"], "category": m["s_category"],
                           "scale": m["s_scale"], "ring": rings.get(sid, [])} for sid, m in meta.items()]

    n = len(raw)
    keep_every = max(1, -(-n // max_points))
    points = []
    prev_stop = None
    for i, r in enumerate(raw):
        f = fit.get(r["id"])
        role = f["role"] if f else None
        method = f["fit"] if f else None
        stop = f["stop"] if f else None
        important = (method == "spike" or role == "reject" or stop != prev_stop
                     or i == 0 or i == n - 1)
        prev_stop = stop
        if not important and i % keep_every:
            continue
        points.append([
            r["dt_message"].isoformat(sep=" "),
            round(float(r["d_lat"]), 6), round(float(r["d_long"]), 6), r["i_speed"],
            round(f["lat"], 6) if f and f["lat"] is not None else None,
            round(f["lon"], 6) if f and f["lon"] is not None else None,
            role, method, None if not f or f["shift_m"] is None else round(f["shift_m"], 1),
            f["reject"] if f else None, stop,
        ])

    routes = []
    for g in gaps:
        geom = polyline.decode(g["s_geometry"], 6) if g["s_geometry"] else [
            (float(g["d_from_lat"]), float(g["d_from_long"])), (float(g["d_to_lat"]), float(g["d_to_long"]))]
        routes.append({"from": g["dt_from"], "to": g["dt_to"], "gap_s": g["i_gap_s"],
                       "routed": g["s_route"] == "ok", "unexplained_s": g["i_unexplained_s"],
                       "path": [[round(a, 6), round(b, 6)] for a, b in geom]})

    lats = [float(r["d_lat"]) for r in raw]
    lons = [float(r["d_long"]) for r in raw]
    return {
        "run_id": rid, "trip_no": trip_no, "fitted": bool(fit),
        "columns": ["ts", "lat", "lon", "speed", "fit_lat", "fit_lon", "role", "fit", "shift_m",
                    "reject", "stop"],
        "total_points": n, "returned": len(points), "points": points,
        "gap_routes": routes, "fences": fences,
        "bbox": [min(lons), min(lats), max(lons), max(lats)],
    }


# ---------------------------------------------------------------------------
# The trip as a number line: loading, transit, unloading (pipeline/phases.py)
# ---------------------------------------------------------------------------

def _report(cur, rid: int, trip_no: int) -> dict:
    t = phases.load_trip(cur, rid, trip_no, phases.run_max_gap(cur, rid))
    if t is None:
        raise HTTPException(404, f"trip {trip_no} is not in run {rid}")
    cur.execute("""SELECT dt_trip_start, dt_booking, dt_trip_ata, s_origin, s_destination
                     FROM geo_trip_meta WHERE i_trip_no=%s""", (trip_no,))
    meta = cur.fetchone() or {}
    segs = [dict(s) for s in t["segments"]]
    for i, s in enumerate(segs):
        s["seq"] = i + 1
    return {"run_id": rid, "trip_no": trip_no, "asset": t["asset"], "start": t["start"], "end": t["end"],
            "summary": t["summary"], "segments": segs,
            "places": [{k: p[k] for k in ("site_id", "name", "scale", "innermost", "enter", "exit",
                                          "last_seen", "open", "dwell_s", "entry_observed")} for p in t["places"]],
            "fleet": {"gate_out": meta.get("dt_trip_start"), "booking": meta.get("dt_booking"),
                      "ata": meta.get("dt_trip_ata"), "origin": meta.get("s_origin"),
                      "destination": meta.get("s_destination")}}


@router.get("/trips/{trip_no}/report")
def trip_report(trip_no: int, run: int | None = None, conn=Depends(get_geo_db)):
    """The whole trip, first fix to last, cut into loading, transit and
    unloading and what filled each -- driving, stops, halts, silent GPS --
    with the time and the kilometre at every boundary."""
    with conn.cursor() as cur:
        rid = resolve_run(cur, run)
        return _report(cur, rid, trip_no)


REPORT_COLUMNS = ("seq", "phase", "kind", "start", "end", "duration_s", "duration_h", "km_from", "km_to", "km",
                  "name", "open", "straight_km")


def _report_rows(rep: dict) -> list[list]:
    out = []
    for s in rep["segments"]:
        km = None if s.get("km_to") is None else round(s["km_to"] - s["km_from"], 2)
        out.append([s["seq"], s["phase"], s["kind"], s["start"], s["end"], s["duration_s"],
                    round(s["duration_s"] / 3600, 2), s.get("km_from"), s.get("km_to"), km,
                    s.get("name"), s.get("open"), s.get("straight_km")])
    return out


def _attachment(name: str) -> dict:
    return {"Content-Disposition": 'attachment; filename="%s"' % name}


@router.get("/trips/{trip_no}/report.csv")
def trip_report_csv(trip_no: int, run: int | None = None, conn=Depends(get_geo_db)):
    import csv
    import io

    with conn.cursor() as cur:
        rid = resolve_run(cur, run)
        rep = _report(cur, rid, trip_no)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(REPORT_COLUMNS)
    for r in _report_rows(rep):
        w.writerow(["" if v is None else v for v in r])
    return Response(buf.getvalue(), media_type="text/csv", headers=_attachment(f"trip-{trip_no}-report.csv"))


@router.get("/trips/{trip_no}/report.xlsx")
def trip_report_xlsx(trip_no: int, run: int | None = None, conn=Depends(get_geo_db)):
    import io

    from openpyxl import Workbook

    with conn.cursor() as cur:
        rid = resolve_run(cur, run)
        rep = _report(cur, rid, trip_no)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Trip", trip_no])
    ws.append(["Vehicle", rep["asset"]])
    ws.append(["GPS from", rep["start"]])
    ws.append(["GPS to", rep["end"]])
    for k, v in rep["summary"].items():
        ws.append([k.replace("_", " "), v])
    seg = wb.create_sheet("Number line")
    seg.append(list(REPORT_COLUMNS))
    for r in _report_rows(rep):
        seg.append(r)
    pl = wb.create_sheet("Places")
    pl.append(["site", "name", "scale", "entered", "left", "hours", "still inside at the end"])
    for p in rep["places"]:
        pl.append([p["site_id"], p["name"], p["scale"], p["enter"], p["exit"], round(p["dwell_s"] / 3600, 2),
                   "yes" if p["open"] else "no"])
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers=_attachment(f"trip-{trip_no}-report.xlsx"))

